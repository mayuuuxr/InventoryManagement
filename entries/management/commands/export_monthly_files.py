"""Save monthly sales/purchase Excel files into Data/YYYY-MM/.

Runs on every app start (current month only - fast) so the owner always
has fresh spreadsheets to work on directly, no export buttons needed.
`--month YYYY-MM` regenerates one month, `--all` rebuilds every month
that has entries (e.g. after an update or import).

Files are views of the database, never the truth: safe to edit, sort
and filter in Excel. Never commit them (gitignored), never ship them
(update-protected).
"""
import calendar as _cal
import datetime as _dt

from django.conf import settings
from django.core.management.base import BaseCommand
from openpyxl import Workbook
from openpyxl.styles import Font

from entries.models import InwardEntry, OutwardEntry

BOLD = Font(bold=True)


def _month_bounds(year, month):
    last = _cal.monthrange(year, month)[1]
    return _dt.date(year, month, 1), _dt.date(year, month, last)


def _months_with_entries():
    months = set()
    for model in (InwardEntry, OutwardEntry):
        for d in model.objects.values_list("date", flat=True).distinct():
            months.add((d.year, d.month))
    return sorted(months)


def _sheet(wb, title, headers, rows, widths):
    ws = wb.active if wb.sheetnames == ["Sheet"] else wb.create_sheet()
    ws.title = title[:31]
    for col, h in enumerate(headers, start=1):
        ws.cell(row=1, column=col, value=h).font = BOLD
    for r, row in enumerate(rows, start=2):
        for col, value in enumerate(row, start=1):
            ws.cell(row=r, column=col, value=value)
    for col, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = w
    return ws


def _totals_sheet(wb, entries):
    """Per-material totals: Material | Unit | Qty. Works for both kinds."""
    totals = {}
    for e in entries:
        key = (e.material.display_name, e.material.unit)
        totals[key] = totals.get(key, 0) + e.qty
    rows = [[name, unit, float(qty)]
            for (name, unit), qty in sorted(totals.items())]
    _sheet(wb, "Totals", ["Material", "Unit", "Qty"], rows, [34, 10, 14])


def export_month(year, month):
    """Write Data/YYYY-MM/{sales,purchases}.xlsx. Returns (dir, n_sales, n_in)."""
    start, end = _month_bounds(year, month)
    out_dir = settings.DATA_DIR / f"{year}-{month:02d}"
    out_dir.mkdir(parents=True, exist_ok=True)

    in_qs = (InwardEntry.objects
             .filter(date__gte=start, date__lte=end)
             .select_related("material").order_by("date", "time", "id"))
    out_qs = (OutwardEntry.objects
              .filter(date__gte=start, date__lte=end)
              .select_related("material").order_by("date", "time", "id"))

    in_rows = list(in_qs.iterator())
    wb = Workbook()
    _sheet(wb, "Purchases",
           ["Date", "Time", "Material", "Qty", "Unit", "Received From"],
           [[e.date.strftime("%Y-%m-%d"),
             e.time.strftime("%I:%M %p").lstrip("0"),
             e.material.display_name, float(e.qty), e.material.unit,
             e.received_from] for e in in_rows],
           [12, 10, 34, 10, 8, 26])
    _totals_sheet(wb, in_rows)
    wb.save(out_dir / "purchases.xlsx")

    out_rows = list(out_qs.iterator())
    wb = Workbook()
    _sheet(wb, "Sales",
           ["Date", "Time", "Material", "Qty", "Unit", "Customer"],
           [[e.date.strftime("%Y-%m-%d"),
             e.time.strftime("%I:%M %p").lstrip("0"),
             e.material.display_name, float(e.qty), e.material.unit,
             e.customer or ""] for e in out_rows],
           [12, 10, 34, 10, 8, 22])
    _totals_sheet(wb, out_rows)
    wb.save(out_dir / "sales.xlsx")
    return out_dir, len(out_rows), len(in_rows)


class Command(BaseCommand):
    help = "Write monthly sales/purchase Excel files into Data/YYYY-MM/."

    def add_arguments(self, parser):
        parser.add_argument("--month", default="",
                            help="YYYY-MM (default: current month)")
        parser.add_argument("--all", action="store_true",
                            help="Rebuild every month that has entries")

    def handle(self, *args, month, all, **options):
        from django.utils import timezone

        if all:
            months = _months_with_entries()
        elif month:
            year, mon = (int(p) for p in month.split("-"))
            months = [(year, mon)]
        else:
            today = timezone.localdate()
            months = [(today.year, today.month)]
        for year, mon in months:
            out_dir, n_sales, n_in = export_month(year, mon)
            self.stdout.write(
                f"{out_dir.name}: {n_sales} sales + {n_in} purchases "
                f"-> {out_dir}")
