"""End-to-end audit suite (permanent regression tests).

Covers: auth/session, role gating, full staff workflows, validation edges,
IDOR, exports, error handlers, query-count performance. No network, no
third-party services: nibba has no photo uploads and no push. Added
during the Sep 2026 pre-production audit; every test here encodes a bug
that was found and fixed (see module docstrings).
Run:  manage.py test entries.tests_audit
"""
from io import BytesIO

from django.test import Client, TestCase, override_settings

from accounts.models import User
from entries.models import InwardEntry, OutwardEntry, SubLocation
from materials.models import Category, Material, SubCategory
from sites.models import Site


def make_world():
    admin = User.objects.create_user(username="a_admin", password="1111",
                                     role=User.Role.ADMIN)
    s1 = User.objects.create_user(username="a_s1", password="2222",
                                  role=User.Role.STAFF)
    s2 = User.objects.create_user(username="a_s2", password="3333",
                                  role=User.Role.STAFF)
    mgr = User.objects.create_user(username="a_mgr", password="4444",
                                   role=User.Role.MANAGER)
    site = Site.objects.create(name="A Site", location="A Loc")
    site2 = Site.objects.create(name="B Site", location="B Loc")
    cat = Category.objects.create(name="A Cat")
    sub = SubCategory.objects.create(category=cat, name="A Sub")
    mat = Material.objects.create(category=cat, sub_category=sub,
                                  name="8mm", unit="nos")
    return admin, s1, s2, mgr, site, site2, cat, sub, mat


@override_settings(ALLOWED_HOSTS=["testserver"])
class AuthAudit(TestCase):
    def setUp(self):
        self.admin, self.s1, _, _, self.site, _, _, _, self.mat = make_world()[:9]

    def test_anon_redirected_to_login(self):
        for url in ["/sites/", "/reports/",
                    f"/entries/sites/{self.site.id}/materials/{self.mat.id}/",
                    "/accounts/staff/", "/entries/activity/"]:
            r = self.client.get(url)
            self.assertEqual(r.status_code, 302, url)
            self.assertIn("/accounts/login/", r["Location"])

    def test_evil_next_rejected(self):
        r = self.client.post("/accounts/login/",
                             {"username": "a_s1", "password": "2222",
                              "next": "https://evil.example/x"})
        self.assertEqual(r.status_code, 302)
        self.assertNotIn("evil.example", r["Location"])

    def test_next_preserved_for_legit(self):
        r = self.client.get("/sites/")
        self.assertIn("next=", r["Location"])

    def test_staff_forbidden_admin_pages(self):
        self.client.force_login(
            User.objects.get(username="a_s1"))
        for url in ["/accounts/staff/", "/accounts/staff/add/",
                    "/entries/activity/", "/entries/activity/all/",
                    "/sites/new/"]:
            r = self.client.get(url)
            self.assertIn(r.status_code, (403, 404, 302), url)

    def test_staff_cannot_delete_site_or_material(self):
        self.client.force_login(User.objects.get(username="a_s1"))
        r = self.client.post(f"/sites/{self.site.id}/delete/")
        self.assertEqual(r.status_code, 403)
        r = self.client.post(f"/materials/{self.mat.id}/delete/")
        self.assertIn(r.status_code, (403, 404))

    def test_logout_clears_session(self):
        self.client.force_login(User.objects.get(username="a_s1"))
        r = self.client.post("/accounts/logout/")
        self.assertEqual(r.status_code, 302)
        r2 = self.client.get("/sites/")
        self.assertEqual(r2.status_code, 302)
        self.assertIn("login", r2["Location"])

    def test_manager_can_write_but_not_activity(self):
        self.client.force_login(User.objects.get(username="a_mgr"))
        r = self.client.get("/entries/activity/")
        self.assertEqual(r.status_code, 403)
        r = self.client.get(
            f"/entries/sites/{self.site.id}/materials/{self.mat.id}/")
        self.assertEqual(r.status_code, 200)


@override_settings(ALLOWED_HOSTS=["testserver"])
class WorkflowAudit(TestCase):
    def setUp(self):
        (self.admin, self.s1, _, _,
         self.site, _, self.cat, _, self.mat) = make_world()
        self.client.force_login(self.s1)
        self.in_url = (f"/entries/sites/{self.site.id}/materials/"
                       f"{self.mat.id}/inward/add/")
        self.out_url = (f"/entries/sites/{self.site.id}/materials/"
                        f"{self.mat.id}/outward/add/")

    def _inward(self, qty=100, t="09:00", src="SRC"):
        return self.client.post(self.in_url, {
            "date": "2026-09-12", "time": t, "qty": str(qty),
            "received_from": src})

    def test_full_inward_outward_balance(self):
        r = self._inward()
        self.assertEqual(r.status_code, 302, r.content[:400])
        inw = InwardEntry.objects.get()
        r = self.client.post(self.out_url, {
            "date": "2026-09-12", "time": "10:00", "qty": "30",
            "sub_location": "Banglo 1", "inward_entry": str(inw.pk)})
        self.assertEqual(r.status_code, 302, r.content[:400])
        from entries.utils import balance_for
        self.assertEqual(balance_for(self.site, self.mat), 70)

    def test_outward_over_balance_blocked(self):
        self._inward(qty=10)
        inw = InwardEntry.objects.get()
        r = self.client.post(self.out_url, {
            "date": "2026-09-12", "time": "10:00", "qty": "11",
            "sub_location": "B1", "inward_entry": str(inw.pk)})
        self.assertEqual(r.status_code, 200)
        self.assertEqual(OutwardEntry.objects.count(), 0)

    def test_negative_decimal_zero_qty_blocked(self):
        for q in ["-5", "10.5", "0"]:
            r = self._inward(qty=q)
            self.assertEqual(r.status_code, 200, q)
        self.assertEqual(InwardEntry.objects.count(), 0)

    def test_entry_saves_with_no_photo_concept(self):
        # nibba has no photo entry at all: plain values save directly.
        r = self.client.post(self.in_url, {
            "date": "2026-09-12", "time": "09:00", "qty": "5",
            "received_from": "X"})
        self.assertEqual(r.status_code, 302, r.content[:400])
        self.assertEqual(InwardEntry.objects.count(), 1)

    def test_duplicate_inward_time_blocked(self):
        self.assertEqual(self._inward(t="09:00").status_code, 302)
        r = self._inward(t="09:00", src="OTHER")
        self.assertEqual(r.status_code, 200)
        self.assertEqual(InwardEntry.objects.count(), 1)

    def test_unicode_long_special_inputs(self):
        uni_src = "üñíçødé ∆ बंगलो " + "x" * 150 + "=CMD|'/ sabotage"
        r = self._inward(qty=7, src=uni_src)
        self.assertEqual(r.status_code, 302, r.content[:400])
        self.assertEqual(InwardEntry.objects.get().received_from, uni_src)
        # Over max_length (200) is rejected, not truncated.
        r = self._inward(qty=7, t="10:00", src="y" * 201)
        self.assertEqual(r.status_code, 200)
        self.assertEqual(InwardEntry.objects.count(), 1)

    def test_huge_qty(self):
        r = self._inward(qty=9999999999)
        self.assertEqual(r.status_code, 302, r.content[:400])

    def test_edit_and_delete_flow(self):
        self._inward(qty=50)
        inw = InwardEntry.objects.get()
        eurl = (f"/entries/sites/{self.site.id}/materials/{self.mat.id}/"
                f"inward/{inw.pk}/edit/")
        r = self.client.post(eurl, {
            "date": "2026-09-12", "time": "09:00", "qty": "60",
            "received_from": "SRC2"})
        self.assertEqual(r.status_code, 302, r.content[:400])
        self.assertEqual(InwardEntry.objects.get().qty, 60)
        durl = (f"/entries/sites/{self.site.id}/materials/{self.mat.id}/"
                f"inward/{inw.pk}/delete/")
        self.assertEqual(self.client.get(durl).status_code, 405)
        r = self.client.post(durl)
        self.assertEqual(r.status_code, 302)
        self.assertEqual(InwardEntry.objects.count(), 0)
        from entries.models import AuditLog
        self.assertTrue(AuditLog.objects.filter(
            action="delete", entry_type="inward").exists())

    def test_delete_blocked_with_links(self):
        self._inward(qty=50)
        inw = InwardEntry.objects.get()
        self.client.post(self.out_url, {
            "date": "2026-09-12", "time": "10:00", "qty": "10",
            "sub_location": "B1", "inward_entry": str(inw.pk)})
        durl = (f"/entries/sites/{self.site.id}/materials/{self.mat.id}/"
                f"inward/{inw.pk}/delete/")
        r = self.client.post(durl, follow=True)
        self.assertEqual(InwardEntry.objects.count(), 1)
        self.assertContains(r, "Cannot delete")

    def test_repeat_prefills_from_any_visible_row(self):
        # Reads are global (owner decision Sep 2026): staff may repeat from
        # a co-worker's row. Edit/delete of others rows stays 404 (IdorAudit).
        self._inward(qty=42, src="SHARED-SRC")
        inw = InwardEntry.objects.get()
        c2 = Client()
        s2 = User.objects.get(username="a_s2")
        c2.force_login(s2)
        r = c2.get(self.in_url + f"?repeat={inw.pk}")
        self.assertEqual(r.status_code, 200)
        form = r.context["form"]
        self.assertEqual(form.initial.get("received_from"), "SHARED-SRC")


@override_settings(ALLOWED_HOSTS=["testserver"])
class IdorAudit(TestCase):
    def setUp(self):
        _, self.s1, self.s2, _, self.site, _, _, _, self.mat = make_world()
        self.client.force_login(self.s1)
        self.inw = InwardEntry.objects.create(
            site=self.site, material=self.mat, date="2026-09-12",
            time="09:00", qty=50, received_from="S1SRC", created_by=self.s1)

    def test_staff_cannot_edit_other_staff_entry(self):
        c2 = Client()
        c2.force_login(self.s2)
        url = (f"/entries/sites/{self.site.id}/materials/{self.mat.id}/"
               f"inward/{self.inw.pk}/edit/")
        self.assertEqual(c2.get(url).status_code, 404)
        r = c2.post(url, {"date": "2026-09-12", "time": "09:00",
                          "qty": "5", "received_from": "HIJACK"})
        self.assertEqual(r.status_code, 404)
        self.assertEqual(InwardEntry.objects.get().qty, 50)

    def test_staff_cannot_delete_other_staff_entry(self):
        c2 = Client()
        c2.force_login(self.s2)
        url = (f"/entries/sites/{self.site.id}/materials/{self.mat.id}/"
               f"inward/{self.inw.pk}/delete/")
        self.assertEqual(c2.post(url).status_code, 404)
        self.assertEqual(InwardEntry.objects.count(), 1)

    def test_staff_timeline_of_other_forbidden(self):
        c2 = Client()
        c2.force_login(self.s2)
        r = c2.get(f"/entries/activity/staff/{self.s1.pk}/")
        self.assertEqual(r.status_code, 403)


@override_settings(ALLOWED_HOSTS=["testserver"])
class ExportAudit(TestCase):
    def setUp(self):
        (self.admin, _, _, _, self.site, _, _, _,
         self.mat) = make_world()
        self.client.force_login(self.admin)
        InwardEntry.objects.create(
            site=self.site, material=self.mat, date="2026-09-12",
            time="09:00", qty=20, received_from="=CMD|' /c calc'!A0",
            created_by=self.admin)

    def test_all_exports_200(self):
        urls = ["/reports/",
                f"/reports/site-summary/?site={self.site.id}",
                f"/reports/site-summary/export/?site={self.site.id}",
                f"/reports/site-summary/export-pdf/?site={self.site.id}",
                f"/reports/material-consumption/?material={self.mat.id}",
                f"/reports/material-consumption/export/?material={self.mat.id}",
                ("/reports/material-consumption/export-pdf/?material="
                 f"{self.mat.id}"),
                "/reports/monthly/?month=9&year=2026",
                "/reports/monthly/export/?month=9&year=2026",
                "/reports/monthly/export-pdf/?month=9&year=2026"]
        for u in urls:
            r = self.client.get(u)
            self.assertEqual(r.status_code, 200, u)

    def test_formula_injection_neutralized(self):
        from openpyxl import load_workbook
        r = self.client.get("/reports/monthly/export/?month=9&year=2026")
        self.assertEqual(r.status_code, 200)
        wb = load_workbook(filename=BytesIO(r.content))
        found = False
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                for v in row:
                    if isinstance(v, str) and "CMD" in v:
                        found = True
                        self.assertTrue(v.startswith("'="), v)
        self.assertTrue(found)

    def test_bad_filters_dont_crash(self):
        for u in ["/reports/monthly/?month=99&year=abcd",
                  "/reports/monthly/export/?month=13&year=2026",
                  "/reports/site-summary/export/",
                  "/reports/material-consumption/export/",
                  f"/entries/sites/9999/materials/{self.mat.id}/",
                  "/entries/sites/abc/materials/xyz/"]:
            r = self.client.get(u)
            self.assertIn(r.status_code, (200, 302, 404), u)

    def test_404_page(self):
        r = self.client.get("/sites/999999/")
        self.assertEqual(r.status_code, 404)

    def test_custom_error_templates_render(self):
        # Test runner forces DEBUG=False, so these exercise the real
        # custom 404/403 templates (which extend base.html + need DB).
        r = self.client.get("/definitely-not-a-page/")
        self.assertEqual(r.status_code, 404)
        self.assertContains(r, "nibba", status_code=404)
        anon = Client()  # never logged in
        r = anon.get("/accounts/staff/")
        self.assertEqual(r.status_code, 302)  # anon -> login
        self.assertIn("/accounts/login/", r["Location"])
        c2 = Client()
        c2.force_login(User.objects.get(username="a_s2"))
        r = c2.get("/entries/activity/")
        self.assertEqual(r.status_code, 403)
        self.assertContains(r, "nibba", status_code=403)


@override_settings(ALLOWED_HOSTS=["testserver"])
class TransferValidationAudit(TestCase):
    """TransferEntry.clean(): positive whole qty, from != to."""

    def setUp(self):
        _, _, _, _, self.site, self.site2, _, _, self.mat = make_world()

    def _t(self, **over):
        from entries.models import TransferEntry
        kw = dict(material=self.mat, date="2026-09-12", time="10:00",
                  qty=10, from_site=self.site, to_site=self.site2)
        kw.update(over)
        return TransferEntry(**kw)

    def test_valid_transfer(self):
        self._t().full_clean()

    def test_same_site_rejected(self):
        from django.core.exceptions import ValidationError
        with self.assertRaises(ValidationError):
            self._t(to_site=self.site).full_clean()

    def test_zero_negative_decimal_rejected(self):
        from django.core.exceptions import ValidationError
        for q in [0, -5, "10.5"]:
            with self.assertRaises(ValidationError, msg=f"qty={q}"):
                self._t(qty=q).full_clean()


@override_settings(ALLOWED_HOSTS=["testserver"])
class AllocationModelGuardAudit(TestCase):
    """Model-level over-allocation block (admin/shell path, not just forms)."""

    def setUp(self):
        (self.admin, _, _, _, self.site, _, _, _,
         self.mat) = make_world()
        self.inw = InwardEntry.objects.create(
            site=self.site, material=self.mat, date="2026-09-12",
            time="09:00", qty=100, received_from="X", created_by=self.admin)
        self.sl = SubLocation.objects.create(site=self.site, name="B1")
        self.sl2 = SubLocation.objects.create(site=self.site, name="B2")

    def test_model_blocks_over_allocation(self):
        from django.core.exceptions import ValidationError
        from entries.models import SubLocationAllocation
        SubLocationAllocation(
            inward_entry=self.inw, sub_location=self.sl,
            qty_allocated=60).save()
        with self.assertRaises(ValidationError):
            SubLocationAllocation(
                inward_entry=self.inw, sub_location=self.sl2,
                qty_allocated=50).save()
        self.assertEqual(SubLocationAllocation.objects.count(), 1)

    def test_edit_self_excluded(self):
        from entries.models import SubLocationAllocation
        a = SubLocationAllocation(
            inward_entry=self.inw, sub_location=self.sl, qty_allocated=60)
        a.save()
        a.qty_allocated = 90  # still within 100: self excluded from siblings
        a.save()
        self.assertEqual(
            SubLocationAllocation.objects.get().qty_allocated, 90)


@override_settings(ALLOWED_HOSTS=["testserver"])
class PerfAudit(TestCase):
    def setUp(self):
        _, _, _, _, self.site, _, self.cat, _, _ = make_world()
        mats = [Material(category=self.cat, name=f"M{i}", unit="nos")
                for i in range(80)]
        Material.objects.bulk_create(mats)
        admin = User.objects.get(username="a_admin")
        self.client.force_login(admin)

    def test_site_detail_query_count(self):
        from django.test.utils import CaptureQueriesContext
        from django.db import connection
        url = f"/sites/{Site.objects.get(name='A Site').id}/"
        with CaptureQueriesContext(connection) as ctx:
            r = self.client.get(url)
        self.assertEqual(r.status_code, 200)
        print(f"\nSITE-DETAIL: {len(ctx)} queries, "
              f"{len(r.content)} bytes")
        self.assertLess(len(ctx), 40)

    def test_material_detail_query_count(self):
        from django.test.utils import CaptureQueriesContext
        from django.db import connection
        mat = Material.objects.first()
        url = (f"/entries/sites/{Site.objects.get(name='A Site').id}/"
               f"materials/{mat.id}/")
        with CaptureQueriesContext(connection) as ctx:
            r = self.client.get(url)
        self.assertEqual(r.status_code, 200)
        print(f"\nMAT-DETAIL: {len(ctx)} queries, "
              f"{len(r.content)} bytes")
        self.assertLess(len(ctx), 60)
