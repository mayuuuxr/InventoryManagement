import json

from django.core.exceptions import ValidationError
from django.test import TestCase
from django.urls import reverse

from accounts.models import User
from materials.models import Category, Material, SubCategory
from sites.models import Site

from .models import InwardEntry, OutwardEntry, SubLocation


class SubLocationFreeTextTests(TestCase):
    """Owner rule: sub-location names are free text — no format checks.
    Anything non-empty saves as-is (trimmed)."""

    def test_anything_nonempty_saves(self):
        site = Site.objects.create(name="Test Site", location="Test")
        for name in ["Banglo 1", "14", "banglo 21 300ft",
                     "Banglo 2, Banglo 4", "Plot_5", "Banglo!"]:
            SubLocation.objects.create(site=site, name=name)
        self.assertEqual(SubLocation.objects.count(), 6)

    def test_empty_names_rejected(self):
        site = Site.objects.create(name="Test Site", location="Test")
        Category.objects.create(name="Test Cat")
        with self.assertRaises(ValidationError):
            SubLocation(site=site, name="   ").save()
        self.assertEqual(SubLocation.objects.count(), 0)


class SingleOutwardFormTests(TestCase):
    """One place per entry: free-text sub-location saves as-is.

    Form-level only (no save → no network, no uploads)."""

    def setUp(self):
        from accounts.models import User
        from materials.models import Material, SubCategory
        from .models import InwardEntry
        self.admin = User.objects.create_user(
            username="s_admin", password="1111", role=User.Role.ADMIN)
        self.site = Site.objects.create(name="S Site", location="S Loc")
        cat = Category.objects.create(name="S Cat")
        sub = SubCategory.objects.create(category=cat, name="S Sub")
        self.material = Material.objects.create(
            category=cat, sub_category=sub, name="S Mat", unit="nos")
        self.inward = InwardEntry.objects.create(
            site=self.site, material=self.material, date="2026-09-12",
            time="09:00", qty=100, received_from="TEST-SRC",
            created_by=self.admin)

    def _form(self, data):
        from .forms import OutwardEntryForm
        return OutwardEntryForm(
            data=data, site=self.site,
            material=self.material, user=self.admin)

    def _base_data(self, **over):
        d = {"date": "2026-09-12", "time": "10:00", "qty": "60",
             "sub_location": "Banglo 2",
             "inward_entry": str(self.inward.pk)}
        d.update(over)
        return d

    def test_valid_single_entry(self):
        form = self._form(self._base_data())
        self.assertTrue(form.is_valid(), form.errors.as_text())
        self.assertEqual(
            form.cleaned_data["sub_location"].name, "Banglo 2")

    def test_free_text_sub_location_accepted(self):
        # Owner rule: sub-location is free text, no format checks —
        # "banglo 21 300ft" saves as-is.
        form = self._form(
            self._base_data(**{"sub_location": "banglo 21 300ft"}))
        self.assertTrue(form.is_valid(), form.errors.as_text())
        self.assertEqual(
            form.cleaned_data["sub_location"].name, "banglo 21 300ft")

    def test_qty_over_delivery_rejected(self):
        form = self._form(self._base_data(**{"qty": "150"}))
        self.assertFalse(form.is_valid())
        self.assertIn("qty", form.errors)


class SingleOutwardViewTests(TestCase):
    """Full POST: one submit saves one entry with free-text sub-location."""

    def setUp(self):
        from accounts.models import User
        from materials.models import Material, SubCategory
        from .models import InwardEntry
        self.admin = User.objects.create_user(
            username="v_admin", password="1111", role=User.Role.ADMIN)
        self.site = Site.objects.create(name="V Site", location="V Loc")
        cat = Category.objects.create(name="V Cat")
        sub = SubCategory.objects.create(category=cat, name="V Sub")
        self.material = Material.objects.create(
            category=cat, sub_category=sub, name="V Mat", unit="nos")
        self.inward = InwardEntry.objects.create(
            site=self.site, material=self.material, date="2026-09-12",
            time="09:00", qty=100, received_from="TEST-SRC",
            created_by=self.admin)
        self.url = (f"/entries/sites/{self.site.id}/materials/"
                    f"{self.material.id}/outward/add/")
        self.client.force_login(self.admin)

    def test_single_place_saved(self):
        from .models import OutwardEntry
        from .utils import balance_for
        r = self.client.post(self.url, {
            "date": "2026-09-12", "time": "10:00", "qty": "60",
            "sub_location": "banglo 21 300ft",
            "inward_entry": str(self.inward.pk)})
        self.assertEqual(r.status_code, 302, r.content[:500])
        rows = list(OutwardEntry.objects.order_by("pk"))
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0].sub_location.name, "banglo 21 300ft")
        self.assertEqual(int(rows[0].qty), 60)
        self.assertEqual(balance_for(self.site, self.material), 40)


class SaleBillViewTests(TestCase):
    """Billing screen: multi-item sale in one bill (entries/sale/bill/)."""

    def setUp(self):
        self.admin = User.objects.create_user(
            username="b_admin", password="1111", role=User.Role.ADMIN)
        self.staff = User.objects.create_user(
            username="b_staff", password="2222", role=User.Role.STAFF)
        # Migration 0003 already created the single "Shop" site — the
        # bill view needs exactly one site to exist (get_default_site).
        self.site = Site.objects.get(name="Shop")
        cat = Category.objects.create(name="B Cat")
        sub = SubCategory.objects.create(category=cat, name="B Sub")
        self.cement = Material.objects.create(
            category=cat, sub_category=sub, name="Cement", unit="bags",
            reorder_point=10)
        self.steel = Material.objects.create(
            category=cat, sub_category=sub, name="Steel 8mm", unit="kg",
            reorder_point=5)
        self.inw = InwardEntry.objects.create(
            site=self.site, material=self.cement, date="2026-10-16",
            time="09:00", qty=100, received_from="SUP", created_by=self.admin)
        self.inw2 = InwardEntry.objects.create(
            site=self.site, material=self.steel, date="2026-10-16",
            time="09:00", qty=50, received_from="SUP", created_by=self.admin)
        self.url = reverse("entries:sale_bill")

    def _post(self, user, payload):
        self.client.force_login(user)
        return self.client.post(
            self.url, data=json.dumps(payload),
            content_type="application/json")

    def test_get_renders_bill_with_catalog(self):
        self.client.force_login(self.admin)
        r = self.client.get(self.url)
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, "New Sale")
        self.assertContains(r, "Type to add product")
        self.assertContains(r, "bill-search")
        self.assertContains(r, "Customer Name")
        body = r.content.decode()
        # Customer field is BLANK — never pre-filled with walk-in text.
        self.assertNotIn('value="Walk-in customer"', body)
        self.assertIn("Cement", body)
        self.assertIn('"balance": 100', body)
        self.assertIn("csrfmiddlewaretoken", body)

    def test_staff_can_save_bill(self):
        r = self._post(self.staff, {
            "customer": "Ramesh", "date": "2026-10-16", "time": "10:00",
            "items": [
                {"material_id": self.cement.pk, "qty": 5},
                {"material_id": self.steel.pk, "qty": 2},
            ]})
        self.assertEqual(r.status_code, 200)
        self.assertTrue(r.json()["ok"])
        # Save lands on the printable thermal bill (Rate/Amount blank),
        # not the orders list.
        self.assertTrue(r.json()["redirect"].startswith(
            "/entries/sale/bill/print/?ids="))
        self.assertEqual(OutwardEntry.objects.count(), 2)
        sale = OutwardEntry.objects.get(material=self.cement)
        self.assertEqual(sale.qty, 5)
        self.assertEqual(sale.customer, "Ramesh")
        self.assertEqual(sale.inward_entry, self.inw)
        self.assertEqual(sale.created_by, self.staff)

    def test_same_material_aggregated_to_one_entry(self):
        r = self._post(self.admin, {
            "customer": "", "date": "2026-10-16", "time": "10:00",
            "items": [
                {"material_id": self.cement.pk, "qty": 3},
                {"material_id": self.cement.pk, "qty": 4},
            ]})
        self.assertEqual(r.status_code, 200)
        self.assertEqual(OutwardEntry.objects.count(), 1)
        self.assertEqual(OutwardEntry.objects.get().qty, 7)

    def test_over_stock_rejected_nothing_saved(self):
        r = self._post(self.admin, {
            "customer": "X", "date": "2026-10-16", "time": "10:00",
            "items": [{"material_id": self.cement.pk, "qty": 101}]})
        self.assertEqual(r.status_code, 400)
        self.assertIn("Only 100 bags", r.json()["error"])
        self.assertEqual(OutwardEntry.objects.count(), 0)

    def test_empty_bill_rejected(self):
        r = self._post(self.admin, {
            "customer": "X", "date": "2026-10-16", "time": "10:00",
            "items": []})
        self.assertEqual(r.status_code, 400)
        self.assertEqual(OutwardEntry.objects.count(), 0)

    def test_duplicate_time_rejected(self):
        OutwardEntry.objects.create(
            site=self.site, material=self.cement, date="2026-10-16",
            time="10:00", qty=1, customer="A", inward_entry=self.inw,
            created_by=self.admin)
        r = self._post(self.admin, {
            "customer": "B", "date": "2026-10-16", "time": "10:00",
            "items": [{"material_id": self.cement.pk, "qty": 1}]})
        self.assertEqual(r.status_code, 400)
        self.assertIn("already has a sale", r.json()["error"])

    def test_anonymous_redirected(self):
        r = self.client.get(self.url)
        self.assertEqual(r.status_code, 302)
        self.assertIn("/accounts/login/", r["Location"])


class PurchaseBillViewTests(TestCase):
    """Purchase billing screen — mirrors the sale bill UI, supplier fields."""

    def setUp(self):
        self.admin = User.objects.create_user(
            username="pb_admin", password="1111", role=User.Role.ADMIN)
        self.staff = User.objects.create_user(
            username="pb_staff", password="2222", role=User.Role.STAFF)
        self.site = Site.objects.get(name="Shop")
        cat = Category.objects.create(name="PB Cat")
        self.paint = Material.objects.create(
            category=cat, name="Acrylic Paint", unit="pcs", reorder_point=5)
        self.tap = Material.objects.create(
            category=cat, name="PVC Tap", unit="nos", reorder_point=10)
        self.url = reverse("entries:purchase_bill")

    def _post(self, user, payload):
        self.client.force_login(user)
        return self.client.post(
            self.url, data=json.dumps(payload),
            content_type="application/json")

    def test_get_renders_bill_with_catalog(self):
        self.client.force_login(self.admin)
        r = self.client.get(self.url)
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, "New Purchase")
        self.assertContains(r, "Supplier Name")
        self.assertContains(r, "Type to add product")
        self.assertContains(r, "bill-search")
        body = r.content.decode()
        self.assertIn("Acrylic Paint", body)

    def test_staff_can_save_purchase_bill(self):
        r = self._post(self.staff, {
            "supplier": "Sharma Paints", "date": "2026-10-17",
            "time": "11:00",
            "items": [
                {"material_id": self.paint.pk, "qty": 10},
                {"material_id": self.tap.pk, "qty": 4},
            ]})
        self.assertEqual(r.status_code, 200)
        self.assertTrue(r.json()["ok"])
        # Straight to the printable purchase bill (thermal, blank prices).
        self.assertTrue(r.json()["redirect"].startswith(
            "/entries/purchase/bill/print/?ids="))
        self.assertEqual(InwardEntry.objects.count(), 2)
        inward = InwardEntry.objects.get(material=self.paint)
        self.assertEqual(inward.qty, 10)
        self.assertEqual(inward.received_from, "Sharma Paints")
        self.assertEqual(inward.created_by, self.staff)

    def test_empty_supplier_rejected(self):
        r = self._post(self.admin, {
            "supplier": "", "date": "2026-10-17", "time": "11:00",
            "items": [{"material_id": self.paint.pk, "qty": 1}]})
        self.assertEqual(r.status_code, 400)
        self.assertIn("supplier name", r.json()["error"].lower())
        self.assertEqual(InwardEntry.objects.count(), 0)

    def test_anonymous_redirected(self):
        r = self.client.get(self.url)
        self.assertEqual(r.status_code, 302)
        self.assertIn("/accounts/login/", r["Location"])


class BillPrintViewTests(TestCase):
    """Printable bills (entries/*bill/print/): 80mm thermal layout,
    Rate/Amount/TOTAL left blank — prices are hand-written on paper."""

    def setUp(self):
        self.admin = User.objects.create_user(
            username="pr_admin", password="1111", role=User.Role.ADMIN)
        self.site = Site.objects.get(name="Shop")
        cat = Category.objects.create(name="Pr Cat")
        self.paint = Material.objects.create(
            category=cat, name="Asian Paint 4L", unit="pcs", reorder_point=5)
        self.out = OutwardEntry.objects.create(
            site=self.site, material=self.paint, date="2026-10-17",
            time="18:30", qty=2, customer="Suresh", created_by=self.admin)
        self.inw = InwardEntry.objects.create(
            site=self.site, material=self.paint, date="2026-10-17",
            time="18:30", qty=5, received_from="Sharma Paints",
            created_by=self.admin)

    def _sale_url(self):
        return reverse("entries:sale_bill_print") + f"?ids={self.out.pk}"

    def _pur_url(self):
        return reverse("entries:purchase_bill_print") + f"?ids={self.inw.pk}"

    def test_sale_print_renders_with_blank_price_cells(self):
        self.client.force_login(self.admin)
        r = self.client.get(self._sale_url())
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, "SALE BILL")
        self.assertContains(r, "Suresh")
        self.assertContains(r, "Asian Paint 4L")
        self.assertContains(r, "TOTAL")
        body = r.content.decode()
        self.assertIn(">Rate</th>", body)
        self.assertIn(">Amt</th>", body)
        # Blank hand-fill cells exist for item amounts and the total.
        self.assertIn('class="blank total-blank"', body)

    def test_sale_print_whole_qty(self):
        self.client.force_login(self.admin)
        body = self.client.get(self._sale_url()).content.decode()
        # Decimal 2.00 renders as plain "2" on the paper.
        self.assertIn(">2 pcs<", body)

    def test_purchase_print_shows_supplier(self):
        self.client.force_login(self.admin)
        r = self.client.get(self._pur_url())
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, "PURCHASE BILL")
        self.assertContains(r, "Sharma Paints")
        self.assertContains(r, "Asian Paint 4L")

    def test_unknown_ids_404(self):
        self.client.force_login(self.admin)
        self.assertEqual(
            self.client.get(
                reverse("entries:sale_bill_print"), {"ids": "999999"}
            ).status_code, 404)
        # Non-numeric junk fails flat, never crashes.
        self.assertEqual(
            self.client.get(
                reverse("entries:sale_bill_print"), {"ids": "abc"}
            ).status_code, 404)
        # Empty ids list = nothing found = 404.
        self.assertEqual(
            self.client.get(
                reverse("entries:sale_bill_print"), {"ids": ""}
            ).status_code, 404)

    def test_anonymous_redirected(self):
        r = self.client.get(reverse("entries:sale_bill_print"))
        self.assertEqual(r.status_code, 302)
        self.assertIn("/accounts/login/", r["Location"])


class OrderDetailViewTests(TestCase):
    """Order sheets opened by clicking a row on the Orders page
    (entries/orders/<inward|outward>/<pk>/): every stored field, the
    sibling lines of the same bill (same party + date + time), and a
    Print receipt link to the 80mm bill paper for the whole bill."""

    def setUp(self):
        self.admin = User.objects.create_user(
            username="od_admin", password="1111", role=User.Role.ADMIN)
        self.site = Site.objects.get(name="Shop")
        cat = Category.objects.create(name="OD Cat")
        self.mat_a = Material.objects.create(
            category=cat, name="Asian Paint 1L", unit="pcs",
            reorder_point=5)
        self.mat_b = Material.objects.create(
            category=cat, name="Putty 5kg", unit="pkts")
        self.mat_o = Material.objects.create(
            category=cat, name="Brush 2in", unit="pcs")
        # One paper bill: same customer + date + time, two lines.
        self.sale_a = OutwardEntry.objects.create(
            site=self.site, material=self.mat_a, date="2026-10-17",
            time="11:30", qty=6, customer="Ramesh Hardware",
            created_by=self.admin)
        self.sale_b = OutwardEntry.objects.create(
            site=self.site, material=self.mat_b, date="2026-10-17",
            time="11:30", qty=2, customer="Ramesh Hardware",
            created_by=self.admin)
        # Different customer at the same date/time — a DIFFERENT bill.
        self.other_sale = OutwardEntry.objects.create(
            site=self.site, material=self.mat_o, date="2026-10-17",
            time="11:30", qty=1, customer="Ganesh Sweets",
            created_by=self.admin)
        self.inw = InwardEntry.objects.create(
            site=self.site, material=self.mat_a, date="2026-10-15",
            time="09:00", qty=50, received_from="Sharma Paints",
            created_by=self.admin)

    def _url(self, kind, pk):
        return reverse("entries:order_detail",
                       kwargs={"kind": kind, "pk": pk})

    def test_requires_login(self):
        for kind, pk in (("outward", self.sale_a.pk),
                         ("inward", self.inw.pk)):
            r = self.client.get(self._url(kind, pk))
            self.assertEqual(r.status_code, 302)
            self.assertIn("/accounts/login/", r["Location"])

    def test_unknown_404s(self):
        self.client.force_login(self.admin)
        self.assertEqual(
            self.client.get(self._url("bogus", 1)).status_code, 404)
        # Outward and Inward are separate tables with separate id
        # sequences: an id that exists nowhere in THIS table is a plain
        # 404 for that URL, never a crash. The inward table holds exactly
        # one row (pk=1), so these ids are unknown to it.
        self.assertEqual(
            self.client.get(
                self._url("inward", self.sale_b.pk)).status_code, 404)
        self.assertEqual(
            self.client.get(
                self._url("inward", 999999)).status_code, 404)
        self.assertEqual(
            self.client.get(
                self._url("outward", 999999)).status_code, 404)

    def test_sale_sheet_shows_fields_and_bill_lines(self):
        self.client.force_login(self.admin)
        r = self.client.get(self._url("outward", self.sale_a.pk))
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, "OUT-" + str(self.sale_a.pk))
        self.assertContains(r, "Ramesh Hardware")
        self.assertContains(r, "Asian Paint 1L")
        self.assertContains(r, "Entered by")
        self.assertContains(r, "od_admin")
        # The sibling line of the SAME bill shows up…
        self.assertContains(r, "Putty 5kg")
        # …the other customer's sale is NOT on this receipt: the bill
        # card shows exactly 2 items. (String-absence can't be asserted
        # for material/customer names — the topbar notification panel
        # echoes the latest rows across the whole app.)
        self.assertContains(r, "Same bill — 2 items")
        # …and the receipt link carries BOTH ids in bill order.
        self.assertContains(
            r, "/entries/sale/bill/print/?ids="
            + f"{self.sale_a.pk},{self.sale_b.pk}")

    def test_fifo_link_when_linked_purchase_exists(self):
        self.client.force_login(self.admin)
        linked = OutwardEntry.objects.create(
            site=self.site, material=self.mat_a, date="2026-10-17",
            time="16:00", qty=3, customer="Deepa", created_by=self.admin,
            inward_entry=self.inw)
        r = self.client.get(self._url("outward", linked.pk))
        self.assertContains(r, "Linked purchase")
        self.assertContains(r, "IN-" + str(self.inw.pk))

    def test_purchase_sheet_shows_supplier_and_receipt_ids(self):
        self.client.force_login(self.admin)
        r = self.client.get(self._url("inward", self.inw.pk))
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, "Sharma Paints")
        self.assertContains(r, "Asian Paint 1L")
        self.assertContains(
            r, "/entries/purchase/bill/print/?ids="
            + str(self.inw.pk))

    def test_print_paper_adds_bill_no(self):
        """The receipt pins the bill to the Orders page: Bill No reuses
        the bill's first line order id (extra Indian-shop detail)."""
        self.client.force_login(self.admin)
        body = self.client.get(
            reverse("entries:sale_bill_print")
            + f"?ids={self.sale_a.pk},{self.sale_b.pk}").content.decode()
        self.assertIn("Bill No:", body)
        self.assertIn("OUT-" + str(self.sale_a.pk), body)
        body = self.client.get(
            reverse("entries:purchase_bill_print")
            + f"?ids={self.inw.pk}").content.decode()
        self.assertIn("IN-" + str(self.inw.pk), body)

    def test_orders_page_rows_link_to_detail(self):
        self.client.force_login(self.admin)
        r = self.client.get(reverse("entries:orders"))
        self.assertContains(r, self._url("outward", self.sale_a.pk))
        self.assertContains(r, self._url("inward", self.inw.pk))


class ExportMonthlyFilesTests(TestCase):
    """export_monthly_files writes Data/YYYY-MM/{sales,purchases}.xlsx
    with entry sheets plus per-material totals."""

    def test_current_month_files(self):
        import tempfile
        from pathlib import Path

        from django.core.management import call_command
        from django.test import override_settings
        from django.utils import timezone
        from accounts.models import User
        from materials.models import Material, SubCategory
        from .models import InwardEntry, OutwardEntry
        admin = User.objects.create_user(
            username="x_admin", password="1111", role=User.Role.ADMIN)
        site = Site.objects.create(name="X Site", location="X Loc")
        cat = Category.objects.create(name="X Cat")
        sub = SubCategory.objects.create(category=cat, name="X Sub")
        mat = Material.objects.create(
            category=cat, sub_category=sub, name="X Mat", unit="nos")
        today = timezone.localdate().isoformat()
        inw = InwardEntry.objects.create(
            site=site, material=mat, date=today, time="09:00:00", qty=100,
            received_from="SUP", created_by=admin)
        OutwardEntry.objects.create(
            site=site, material=mat, date=today, time="10:00:00", qty=30,
            customer="Buyer", inward_entry=inw, created_by=admin)
        with tempfile.TemporaryDirectory() as tmp:
            with override_settings(DATA_DIR=Path(tmp)):
                call_command("export_monthly_files")
            tag = timezone.localdate().strftime("%Y-%m")
            sales = Path(tmp) / tag / "sales.xlsx"
            purch = Path(tmp) / tag / "purchases.xlsx"
            self.assertTrue(sales.exists())
            self.assertTrue(purch.exists())
            from openpyxl import load_workbook
            ws = load_workbook(sales)["Sales"]
            self.assertEqual(ws["A1"].value, "Date")
            self.assertEqual(ws["D2"].value, 30)
            self.assertEqual(ws["F2"].value, "Buyer")
            tot = load_workbook(sales)["Totals"]
            self.assertEqual(tot["C2"].value, 30)
            wp = load_workbook(purch)["Purchases"]
            self.assertEqual(wp["D2"].value, 100)
